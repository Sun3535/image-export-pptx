"""Extract every picture in a .pptx as a separate lossless TIFF."""

from __future__ import annotations

import hashlib
import json
from collections.abc import Iterator
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
from pptx import Presentation

from . import __version__
from .models import STATUS_ERROR, STATUS_OK, STATUS_SKIPPED_VECTOR, ImageRecord
from .tiff import VECTOR_EXTS, DecodedImage, decode, save_tiff
from .walker import PictureRef, iter_pictures

EMU_PER_INCH = 914400


def _stem(slide_index: int, image_index: int) -> str:
    return f"slide{slide_index:03d}_img{image_index:02d}"


def _effective_ppi(rec: ImageRecord) -> float | None:
    """Pixels per inch of the visible (cropped) part as displayed on the slide."""
    w_emu, h_emu = rec.bbox_emu.get("width", 0), rec.bbox_emu.get("height", 0)
    if not (w_emu and h_emu and rec.width_px and rec.height_px):
        return None
    c = rec.crop
    vis_w = rec.width_px * (1 - c["left"] - c["right"])
    vis_h = rec.height_px * (1 - c["top"] - c["bottom"])
    ppi = min(vis_w / (w_emu / EMU_PER_INCH), vis_h / (h_emu / EMU_PER_INCH))
    return round(ppi, 1)


def _add_message(rec: ImageRecord, text: str | None) -> None:
    if text:
        rec.message = f"{rec.message}; {text}" if rec.message else text


def _process(pptx_path: Path) -> Iterator[tuple[ImageRecord, DecodedImage | None, bytes | None, str | None]]:
    """Yield (record, decoded image, original bytes, original extension) for every picture."""
    prs = Presentation(str(pptx_path))
    for slide_index, slide in enumerate(prs.slides, start=1):
        refs: list[PictureRef] = iter_pictures(slide)
        for image_index, ref in enumerate(refs, start=1):
            shape = ref.shape
            rec = ImageRecord(
                slide_index=slide_index,
                image_index=image_index,
                shape_id=shape.shape_id,
                shape_name=shape.name,
                group_path=ref.group_path,
                crop={
                    "left": shape.crop_left,
                    "top": shape.crop_top,
                    "right": shape.crop_right,
                    "bottom": shape.crop_bottom,
                },
                bbox_emu=ref.bbox_emu,
            )
            try:
                image = shape.image
                blob, ext = image.blob, image.ext.lower()
            except Exception as exc:  # e.g. linked (not embedded) picture
                rec.status = STATUS_ERROR
                rec.message = f"cannot read embedded image: {exc}"
                yield rec, None, None, None
                continue

            rec.sha256 = hashlib.sha256(blob).hexdigest()
            rec.source_format = ext
            if ext in VECTOR_EXTS:
                rec.status = STATUS_SKIPPED_VECTOR
                rec.message = f"vector image ({ext}) is not converted to TIFF; original saved"
                yield rec, None, blob, ext
                continue

            try:
                decoded = decode(blob)
            except Exception as exc:
                rec.status = STATUS_ERROR
                rec.message = f"cannot decode image: {exc}"
                yield rec, None, blob, ext
                continue

            rec.width_px, rec.height_px = decoded.width, decoded.height
            rec.mode = decoded.mode
            rec.bits_per_sample = decoded.bits_per_sample
            rec.dpi = decoded.dpi
            rec.effective_ppi = _effective_ppi(rec)
            _add_message(rec, decoded.note)
            yield rec, decoded, blob, ext


def iter_images(pptx_path: str | Path) -> Iterator[tuple[ImageRecord, np.ndarray]]:
    """Yield (record, pixel array) for every raster picture, without writing files.

    The array keeps the full resolution and bit depth of the embedded image,
    so it can go straight into segmentation code.
    """
    for rec, decoded, _, _ in _process(Path(pptx_path)):
        if decoded is not None:
            yield rec, decoded.to_array()


def extract_images(
    pptx_path: str | Path,
    out_dir: str | Path | None = None,
    compression: str = "deflate",
    keep_original: bool = False,
    min_ppi: float | None = None,
) -> list[ImageRecord]:
    """Save every picture in ``pptx_path`` as its own TIFF under ``out_dir/<pptx name>/``.

    ``out_dir`` defaults to the folder the .pptx is in.
    Also writes ``manifest.json`` describing each image. Returns the records.
    """
    pptx_path = Path(pptx_path)
    target = Path(out_dir if out_dir is not None else pptx_path.parent) / pptx_path.stem
    target.mkdir(parents=True, exist_ok=True)

    records: list[ImageRecord] = []
    slide_count = len(Presentation(str(pptx_path)).slides)
    for rec, decoded, blob, ext in _process(pptx_path):
        stem = _stem(rec.slide_index, rec.image_index)
        if blob is not None and (keep_original or rec.status == STATUS_SKIPPED_VECTOR):
            orig = target / f"{stem}.orig.{ext}"
            orig.write_bytes(blob)
            rec.original_path = orig.name
        if decoded is not None:
            tif = target / f"{stem}.tif"
            try:
                _add_message(rec, save_tiff(decoded, tif, compression))
                rec.output_path = tif.name
            except Exception as exc:
                rec.status = STATUS_ERROR
                _add_message(rec, f"cannot write TIFF: {exc}")
        if (
            rec.status == STATUS_OK
            and min_ppi is not None
            and rec.effective_ppi is not None
            and rec.effective_ppi < min_ppi
        ):
            _add_message(
                rec,
                f"low on-slide resolution: {rec.effective_ppi} ppi (< {min_ppi}); "
                "saved pixels are unchanged, but check the source was not "
                "downsized by PowerPoint image compression",
            )
        records.append(rec)

    manifest = {
        "source_file": pptx_path.name,
        "slide_count": slide_count,
        "image_count": len(records),
        "compression": compression,
        "tool_version": __version__,
        "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "images": [r.to_dict() for r in records],
    }
    (target / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return records

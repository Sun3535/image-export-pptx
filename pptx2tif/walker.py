"""Find every picture shape on a slide, including those nested in groups."""

from __future__ import annotations

from dataclasses import dataclass, field

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.shapes.picture import Picture

# Affine map from a shape's own coordinate space to slide coordinates:
# slide_x = sx * x + tx, slide_y = sy * y + ty
_IDENTITY = (1.0, 0.0, 1.0, 0.0)


@dataclass
class PictureRef:
    shape: Picture
    group_path: list[str] = field(default_factory=list)
    bbox_emu: dict[str, int] = field(default_factory=dict)


def _apply(tf, left, top, width, height) -> dict[str, int]:
    sx, tx, sy, ty = tf
    return {
        "left": round(sx * left + tx),
        "top": round(sy * top + ty),
        "width": round(sx * width),
        "height": round(sy * height),
    }


def _group_transform(group, parent_tf):
    """Compose the parent transform with the group's child-space mapping."""
    xfrm = group._element.grpSpPr.find(qn("a:xfrm"))
    if xfrm is None:
        return parent_tf
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    ch_off, ch_ext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
    if None in (off, ext, ch_off, ch_ext):
        return parent_tf

    def scale(full, child):
        child = int(child)
        return int(full) / child if child else 1.0

    gsx = scale(ext.get("cx"), ch_ext.get("cx"))
    gsy = scale(ext.get("cy"), ch_ext.get("cy"))
    gtx = int(off.get("x")) - gsx * int(ch_off.get("x"))
    gty = int(off.get("y")) - gsy * int(ch_off.get("y"))
    psx, ptx, psy, pty = parent_tf
    return (psx * gsx, psx * gtx + ptx, psy * gsy, psy * gty + pty)


def _walk(shapes, group_path, tf, out):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            _walk(
                shape.shapes,
                group_path + [shape.name],
                _group_transform(shape, tf),
                out,
            )
        elif isinstance(shape, Picture):  # also covers PlaceholderPicture
            bbox = _apply(
                tf,
                shape.left or 0,
                shape.top or 0,
                shape.width or 0,
                shape.height or 0,
            )
            out.append(PictureRef(shape, list(group_path), bbox))


def iter_pictures(slide) -> list[PictureRef]:
    """Return the slide's pictures in reading order (top-to-bottom, then left-to-right)."""
    found: list[PictureRef] = []
    _walk(slide.shapes, [], _IDENTITY, found)
    found.sort(key=lambda p: (p.bbox_emu["top"], p.bbox_emu["left"]))
    return found

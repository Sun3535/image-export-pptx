"""Build test .pptx files in code so the tests need no binary fixtures."""

from __future__ import annotations

import io

import numpy as np
import png
import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

RNG = np.random.default_rng(0)


def png_rgb8(w=64, h=48) -> bytes:
    buf = io.BytesIO()
    Image.fromarray(RNG.integers(0, 256, (h, w, 3), dtype=np.uint8), "RGB").save(buf, "PNG")
    return buf.getvalue()


def png_rgba8(w=50, h=40) -> bytes:
    buf = io.BytesIO()
    Image.fromarray(RNG.integers(0, 256, (h, w, 4), dtype=np.uint8), "RGBA").save(buf, "PNG")
    return buf.getvalue()


def jpeg_rgb(w=80, h=60) -> bytes:
    buf = io.BytesIO()
    Image.fromarray(RNG.integers(0, 256, (h, w, 3), dtype=np.uint8), "RGB").save(
        buf, "JPEG", quality=90
    )
    return buf.getvalue()


def png_gray16(w=70, h=30) -> bytes:
    arr = RNG.integers(0, 65536, (h, w), dtype=np.uint16)
    buf = io.BytesIO()
    png.Writer(w, h, greyscale=True, bitdepth=16).write(buf, arr.tolist())
    return buf.getvalue()


def png_rgb16(w=40, h=20) -> bytes:
    arr = RNG.integers(0, 65536, (h, w, 3), dtype=np.uint16)
    buf = io.BytesIO()
    png.Writer(w, h, greyscale=False, bitdepth=16).write(buf, arr.reshape(h, -1).tolist())
    return buf.getvalue()


def png_palette(w=30, h=30) -> bytes:
    img = Image.fromarray(RNG.integers(0, 16, (h, w), dtype=np.uint8), "L").convert("P")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


@pytest.fixture
def sample_pptx(tmp_path):
    """A deck covering the cases the extractor must handle.

    Returns (path, expected) where expected maps output stem -> source bytes.
    """
    prs = Presentation()
    blank = prs.slide_layouts[6]
    expected: dict[str, bytes] = {}

    def add(slide, blob, left, top, width=None, height=None, shapes=None):
        target = slide.shapes if shapes is None else shapes  # empty group is falsy
        return target.add_picture(
            io.BytesIO(blob), Inches(left), Inches(top),
            Inches(width) if width else None, Inches(height) if height else None,
        )

    # Slide 1: a single image.
    s1 = prs.slides.add_slide(blank)
    b = png_rgb8()
    add(s1, b, 1, 1)
    expected["slide001_img01"] = b

    # Slide 2: three images, added out of reading order.
    s2 = prs.slides.add_slide(blank)
    rgba, jpg, g16 = png_rgba8(), jpeg_rgb(), png_gray16()
    add(s2, g16, 1, 5)       # bottom
    add(s2, jpg, 6, 1)       # top right
    add(s2, rgba, 1, 1)      # top left
    expected["slide002_img01"] = rgba
    expected["slide002_img02"] = jpg
    expected["slide002_img03"] = g16

    # Slide 3: two images inside a group, one 16-bit RGB; plus a cropped image.
    s3 = prs.slides.add_slide(blank)
    grp = s3.shapes.add_group_shape()
    rgb16, pal = png_rgb16(), png_palette()
    add(s3, rgb16, 1, 1, shapes=grp.shapes)
    add(s3, pal, 4, 1, shapes=grp.shapes)
    cropped_blob = png_rgb8(100, 100)
    cropped = add(s3, cropped_blob, 1, 4)
    cropped.crop_left, cropped.crop_right = 0.25, 0.1
    expected["slide003_img01"] = rgb16
    expected["slide003_img02"] = pal
    expected["slide003_img03"] = cropped_blob

    # Slide 4: the same image used twice must give two files.
    s4 = prs.slides.add_slide(blank)
    dup = png_rgb8(20, 20)
    add(s4, dup, 1, 1, 4, 4)  # 20 px over 4 inch -> 5 ppi, triggers low-ppi warning
    add(s4, dup, 6, 1)
    expected["slide004_img01"] = dup
    expected["slide004_img02"] = dup

    # Slide 5: no images.
    prs.slides.add_slide(blank)

    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path, expected
